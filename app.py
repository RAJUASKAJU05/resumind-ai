"""Resumind AI - compact full-stack app (FastAPI + SQLite). Run: uvicorn app:app"""
import os, re, io, json, hmac, time, base64, hashlib, secrets, sqlite3, unicodedata
from pathlib import Path
from fastapi import FastAPI, UploadFile, File, HTTPException, Request, Response, Depends
from fastapi.responses import FileResponse
from pydantic import BaseModel, EmailStr, Field

DB = os.getenv("DB_PATH", "resumind.db")
SECRET = os.getenv("JWT_SECRET") or secrets.token_hex(32)  # set JWT_SECRET to keep sessions across restarts
AI_KEY = os.getenv("AI_API_KEY", "")
AI_MODEL = os.getenv("AI_MODEL", "gemini-2.0-flash")
MAX_MB = 5
app = FastAPI(title="Resumind AI")

@app.exception_handler(Exception)
async def server_error(request: Request, exc: Exception):
    import logging; from fastapi.responses import JSONResponse
    logging.getLogger("resumind").exception("unhandled error on %s", request.url.path)  # details stay in server logs only
    return JSONResponse({"detail": "Something went wrong on our side. Please try again in a moment."}, status_code=500)

DATABASE_URL = os.getenv("DATABASE_URL", "")
PG = DATABASE_URL.startswith("postgres")

class Res:
    """Cursor wrapper that keeps its connection alive until the result is used."""
    def __init__(self, conn, cur): self.conn, self.cur = conn, cur
    def fetchone(self): return self.cur.fetchone()
    def fetchall(self): return self.cur.fetchall()
    def __iter__(self): return iter(self.cur.fetchall())
    @property
    def rowcount(self): return self.cur.rowcount
    @property
    def lastrowid(self): return getattr(self, "_id", None) or getattr(self.cur, "lastrowid", None)

class Conn:
    """Small wrapper so the same SQL runs on SQLite (local) and Postgres/Supabase (production)."""
    def __init__(self):
        if PG:
            import psycopg
            from psycopg.rows import dict_row
            self.c = psycopg.connect(DATABASE_URL, row_factory=dict_row, connect_timeout=10, prepare_threshold=None)
        else:
            self.c = sqlite3.connect(DB); self.c.row_factory = sqlite3.Row; self.c.execute("PRAGMA foreign_keys=ON")
    def execute(self, q, a=()):
        if not PG: return Res(self, self.c.execute(q, a))
        q = q.replace("?", "%s"); ins = q.lstrip().upper().startswith("INSERT") and "ON CONFLICT" not in q.upper()
        cur = self.c.cursor(); cur.execute(q + (" RETURNING id" if ins else ""), a)
        res = Res(self, cur)
        if ins: r = cur.fetchone(); res._id = r["id"] if r else None
        return res
    def executemany(self, q, rows):
        cur = self.c.cursor() if PG else self.c
        cur.executemany(q.replace("?", "%s") if PG else q, rows); return Res(self, cur)
    def commit(self): self.c.commit()
    def close(self):
        try: self.c.close()
        except Exception: pass
    def __del__(self): self.close()

def db(): return Conn()

def init():
    if PG: return  # schema is managed by the Supabase migration (see supabase_schema.sql)
    c = db(); c.c.executescript("""
    CREATE TABLE IF NOT EXISTS users(id INTEGER PRIMARY KEY, name TEXT, email TEXT UNIQUE, pw TEXT, created TEXT DEFAULT CURRENT_TIMESTAMP);
    CREATE TABLE IF NOT EXISTS resumes(id INTEGER PRIMARY KEY, user_id INTEGER REFERENCES users(id) ON DELETE CASCADE,
      filename TEXT, text TEXT, parsed TEXT, analysis TEXT, created TEXT DEFAULT CURRENT_TIMESTAMP);
    CREATE TABLE IF NOT EXISTS jobs(id INTEGER PRIMARY KEY, user_id INTEGER REFERENCES users(id) ON DELETE CASCADE,
      title TEXT, description TEXT, created TEXT DEFAULT CURRENT_TIMESTAMP);
    CREATE TABLE IF NOT EXISTS matches(id INTEGER PRIMARY KEY, resume_id INTEGER REFERENCES resumes(id) ON DELETE CASCADE,
      job_id INTEGER REFERENCES jobs(id) ON DELETE CASCADE, score INTEGER, result TEXT, created TEXT DEFAULT CURRENT_TIMESTAMP);
    CREATE TABLE IF NOT EXISTS messages(id INTEGER PRIMARY KEY, user_id INTEGER REFERENCES users(id) ON DELETE CASCADE, role TEXT, content TEXT, created TEXT DEFAULT CURRENT_TIMESTAMP);
    CREATE TABLE IF NOT EXISTS applications(id INTEGER PRIMARY KEY, user_id INTEGER REFERENCES users(id) ON DELETE CASCADE, company TEXT, role TEXT,
      status TEXT DEFAULT 'Saved', notes TEXT DEFAULT '', resume_id INTEGER REFERENCES resumes(id) ON DELETE SET NULL, job_id INTEGER REFERENCES jobs(id) ON DELETE SET NULL, created TEXT DEFAULT CURRENT_TIMESTAMP);
    CREATE TABLE IF NOT EXISTS interviews(id INTEGER PRIMARY KEY, user_id INTEGER REFERENCES users(id) ON DELETE CASCADE, question TEXT, answer TEXT, score INTEGER, feedback TEXT, created TEXT DEFAULT CURRENT_TIMESTAMP);
    CREATE INDEX IF NOT EXISTS ix_res_user ON resumes(user_id);""")
    c.commit(); c.close()
init()

# ---------- auth ----------
_demo = (os.getenv('DEMO_EMAIL'), os.getenv('DEMO_PASSWORD'))
def hash_pw(p, salt=None):
    salt = salt or secrets.token_hex(8)
    return salt + "$" + hashlib.pbkdf2_hmac("sha256", p.encode(), salt.encode(), 200_000).hex()
def check_pw(p, h):
    return hmac.compare_digest(hash_pw(p, h.split("$")[0]), h)
def make_token(uid):
    body = base64.urlsafe_b64encode(json.dumps({"u": uid, "e": int(time.time()) + 7*86400}).encode()).decode()
    return body + "." + hmac.new(SECRET.encode(), body.encode(), hashlib.sha256).hexdigest()
def user(request: Request):
    t = request.cookies.get("session", "")
    try:
        body, sig = t.split(".")
        if not hmac.compare_digest(sig, hmac.new(SECRET.encode(), body.encode(), hashlib.sha256).hexdigest()): raise ValueError
        d = json.loads(base64.urlsafe_b64decode(body))
        if d["e"] < time.time(): raise ValueError
        return d["u"]
    except Exception:
        raise HTTPException(401, "Please log in.")

_hits = {}
def limit(key, n=10, per=60):
    now = time.time(); h = [t for t in _hits.get(key, []) if now - t < per]
    if len(h) >= n: raise HTTPException(429, "Too many attempts. Please wait a minute.")
    _hits[key] = h + [now]

class Reg(BaseModel):
    name: str = Field(min_length=1, max_length=80); email: EmailStr; password: str = Field(min_length=8, max_length=128)
class Login(BaseModel):
    email: str; password: str

def set_cookie(resp, uid):
    resp.set_cookie("session", make_token(uid), httponly=True, samesite="lax", max_age=7*86400, secure=os.getenv("COOKIE_SECURE") == "1")

def seed_demo():
    """Free hosting wipes the database on restart; this re-creates a shared demo login at every start."""
    if all(_demo):
        c = db(); c.execute("INSERT INTO users(name,email,pw) VALUES(?,?,?) ON CONFLICT(email) DO NOTHING", ("Demo", _demo[0].lower(), hash_pw(_demo[1]))); c.commit(); c.close()
seed_demo()

@app.post("/api/auth/register", status_code=201)
def register(b: Reg, request: Request, resp: Response):
    limit("r" + request.client.host)
    c = db()
    try:
        cur = c.execute("INSERT INTO users(name,email,pw) VALUES(?,?,?)", (b.name.strip(), b.email.lower(), hash_pw(b.password))); c.commit()
    except Exception as e:
        if "unique" not in str(e).lower() and "integrity" not in type(e).__name__.lower(): raise
        raise HTTPException(409, "An account with this email already exists.")
    set_cookie(resp, cur.lastrowid); return {"id": cur.lastrowid, "name": b.name}

@app.post("/api/auth/login")
def login(b: Login, request: Request, resp: Response):
    limit("l" + request.client.host)
    r = db().execute("SELECT * FROM users WHERE email=?", (b.email.lower(),)).fetchone()
    if not r or not check_pw(b.password, r["pw"]): raise HTTPException(401, "Incorrect email or password.")
    set_cookie(resp, r["id"]); return {"id": r["id"], "name": r["name"]}

@app.post("/api/auth/logout")
def logout(resp: Response): resp.delete_cookie("session"); return {"ok": True}

@app.get("/api/auth/me")
def me(uid=Depends(user)):
    r = db().execute("SELECT id,name,email FROM users WHERE id=?", (uid,)).fetchone()
    if not r: raise HTTPException(401, "Please log in.")
    return dict(r)

# ---------- parsing ----------
SKILLS = """python java javascript typescript c++ c# golang rust php ruby kotlin swift sql mysql postgresql mongodb redis sqlite html css react angular vue nextjs node expressjs
django flask fastapi spring-boot docker kubernetes aws azure gcp git github linux bash rest-api api graphql tensorflow pytorch pandas numpy scikit-learn nlp machine-learning
deep-learning data-analysis excel power-bi tableau plc scada arduino raspberry-pi embedded iot matlab autocad jenkins ci/cd agile scrum testing selenium jira figma communication
teamwork leadership problem-solving terraform ansible prometheus grafana nginx rabbitmq kafka spark hadoop airflow firebase supabase tailwind bootstrap flutter android opencv""".split()
# Written forms -> canonical skill. Ambiguous English words (go, rest, express, spring, c) only count in these specific forms.
ALIASES = {"machine learning": "machine-learning", "deep learning": "deep-learning", "data analysis": "data-analysis", "power bi": "power-bi", "scikit learn": "scikit-learn",
           "node.js": "node", "nodejs": "node", "next.js": "nextjs", "react.js": "react", "reactjs": "react", "express.js": "expressjs", "raspberry pi": "raspberry-pi",
           "problem solving": "problem-solving", "restful": "rest-api", "rest api": "rest-api", "rest apis": "rest-api", "spring boot": "spring-boot", "golang": "golang", "go lang": "golang"}
HEADINGS = {"summary": r"summary|objective|profile|about me", "skills": r"skills|technologies|tech stack|competenc\w+", "education": r"education|academic|qualifications?",
            "experience": r"experience|employment|work history|internships?", "projects": r"projects?", "certifications": r"certifications?|certificates|courses|licenses",
            "achievements": r"achievements|awards|honou?rs|accomplishments"}
JOB_WORDS = {"manager", "engineer", "developer", "analyst", "lead", "intern", "officer", "designer", "consultant", "administrator", "architect", "researcher"}
BULLET = "•▪●◦‣-*–"

def clean(text):
    text = text.replace("\xad", "").replace("\u200b", "").replace("\xa0", " ")
    return re.sub(r"[ \t]+", " ", text).strip()

def heading_of(line):
    """Return the section key if this short line looks like a section heading (e.g. 'Work Experience', 'TECHNICAL SKILLS', 'Honors & Awards')."""
    l = line.strip().rstrip(":").strip(); words = l.lower().replace("&", " ").split()
    if not l or len(words) > 4 or l[0] in BULLET or l.endswith((".", ",")) or JOB_WORDS & set(words): return None
    if not (l.isupper() or l.istitle() or all(w[0].isupper() or w in ("and", "&", "of") for w in l.split())): return None
    return next((k for k, p in HEADINGS.items() if re.search(rf"\b({p})\b", l.lower())), None)

def find_skills(text):
    t = " " + text.lower() + " "
    for a in sorted(ALIASES, key=len, reverse=True): t = t.replace(a, ALIASES[a])  # longest first so 'rest apis' wins over 'rest api'
    return sorted({s for s in SKILLS if re.search(r"(?<![\w+#/-])" + re.escape(s) + r"(?![\w+#-])", t)})

def _tables(doc):
    try: return sum(len(p.find_tables().tables) for p in doc)
    except Exception: return 0

def ocr_image(img_bytes):
    """OCR with preprocessing (upscale, grayscale, contrast) and several page-layout modes; keeps the best result."""
    try:
        import pytesseract
        from PIL import Image, ImageOps, ImageFilter
        img = ImageOps.exif_transpose(Image.open(io.BytesIO(img_bytes))).convert("L")
        if img.width < 2000: img = img.resize((2000, int(img.height * 2000 / img.width)), Image.LANCZOS)
        img = ImageOps.autocontrast(img).filter(ImageFilter.SHARPEN)
        best = ""
        for psm in (6, 4, 3):
            t = pytesseract.image_to_string(img, config=f"--psm {psm}")
            if len(re.findall(r"\w", t)) > len(re.findall(r"\w", best)): best = t
        return best
    except Exception:
        raise HTTPException(422, "We couldn't read text from this image. Try a clearer, higher-resolution scan or upload the PDF/DOCX.")

def extract_text(name, data):
    ext = Path(name).suffix.lower()
    if ext in (".png", ".jpg", ".jpeg"):
        if not (data.startswith(b"\x89PNG") or data.startswith(b"\xff\xd8")): raise HTTPException(400, "Please upload a PDF, DOCX, PNG or JPG file.")
        return clean(ocr_image(data)), {"images": 0, "tables": 0}
    if ext == ".pdf":
        if not data.startswith(b"%PDF"): raise HTTPException(400, "Please upload a PDF, DOCX, PNG or JPG file.")
        import fitz
        try: doc = fitz.open(stream=data, filetype="pdf")
        except Exception: raise HTTPException(422, "We couldn't read this PDF.")
        if doc.needs_pass: raise HTTPException(422, "This PDF is password-protected. Please remove the password and upload it again.")
        text = "\n".join(p.get_text() for p in doc)
        meta = {"images": sum(len(p.get_images()) for p in doc), "tables": _tables(doc)}
        if len(text.strip()) < 100:  # scanned PDF: OCR the first pages
            text = "\n".join(ocr_image(p.get_pixmap(dpi=200).tobytes("png")) for p in list(doc)[:3]); meta["images"] = 0
        return clean(text), meta
    if ext == ".docx":
        if not data.startswith(b"PK"): raise HTTPException(400, "Please upload a PDF, DOCX, PNG or JPG file.")
        import docx
        try: d = docx.Document(io.BytesIO(data))
        except Exception: raise HTTPException(422, "We couldn't read this DOCX.")
        cells = [c.text for t in d.tables for row in t.rows for c in row.cells]  # text inside tables is part of the resume too
        return clean("\n".join([p.text for p in d.paragraphs] + cells)), {"images": len(d.inline_shapes), "tables": len(d.tables)}
    raise HTTPException(400, "Please upload a PDF, DOCX, PNG or JPG file.")

def parse(text):
    lines = [l.strip() for l in text.splitlines() if l.strip()]
    sections, cur = {}, "header"
    for l in lines:
        h = heading_of(l)
        if h: cur = h; sections.setdefault(cur, []); continue
        sections.setdefault(cur, []).append(l)
    g = lambda p: (re.search(p, text, re.I) or [None])[0]
    name = next((l for l in lines[:5] if not re.search(r"resume|curriculum|\bcv\b|@|\d", l, re.I)), lines[0] if lines else "")
    return {"name": name, "email": g(r"[\w.+-]+@[\w-]+\.[\w.]+"), "phone": g(r"(?:\+|\()?\+?\d[\d\s().-]{8,}\d"),
            "linkedin": g(r"linkedin\.com/[\w/%-]+"), "github": g(r"github\.com/[\w-]+"),
            "has_profile_link": bool(re.search(r"(linkedin|github)\s*(\.com/|:)", text, re.I)),
            "sections": {k: v for k, v in sections.items() if k != "header"}, "skills": find_skills(text)}

# ---------- deterministic scoring + ATS ----------
IRREGULAR = {"built", "led", "ran", "wrote", "made", "drove", "grew", "won", "taught", "began", "set", "put", "cut", "chose", "spoke"}
BASE_VERBS = {"build", "develop", "design", "implement", "create", "lead", "manage", "improve", "reduce", "increase", "automate", "deploy", "analyze", "optimize", "integrate", "maintain", "test", "configure", "write", "support", "monitor", "mentor", "architect"}
def is_action(word):
    w = word.lower().strip(BULLET + " ,.:;")
    return w in IRREGULAR or w in BASE_VERBS or (w.endswith("ed") and len(w) > 4)

def bullet_items(lines):
    """Items that start with a bullet marker; if the section uses none, every reasonably long line counts."""
    items, cur = [], None
    for l in lines:
        if l[0] in BULLET:
            if cur: items.append(cur)
            cur = l.lstrip(BULLET + " ")
        elif cur is not None and l[0].islower(): cur += " " + l  # wrapped continuation of the previous bullet
        else:
            if cur: items.append(cur)
            cur = None
    if cur: items.append(cur)
    return [i for i in items if len(i.split()) > 3] or [l for l in lines if len(l.split()) > 4]

def analyse(text, parsed, meta):
    s = parsed["sections"]; words = len(text.split())
    bullets = [b for k in ("experience", "projects") for b in bullet_items(s.get(k, []))]
    verb_hits = sum(1 for b in bullets if is_action(b.split()[0]))
    metrics = sum(1 for b in bullets if re.search(r"\d+\s*%|\d{2,}|\$\d", b))
    checks, recs = [], []
    def chk(ok, label, fix, w):
        checks.append({"label": label, "pass": bool(ok), "weight": w})
        if not ok: recs.append(fix)
    chk(parsed["email"], "Email found", "Add a professional email address in the header.", 10)
    chk(parsed["phone"], "Phone found", "Add a phone number.", 5)
    chk(parsed["linkedin"] or parsed["github"] or parsed["has_profile_link"], "LinkedIn or GitHub link", "Add a LinkedIn or GitHub profile link.", 5)
    for k, lab in (("education", "Education"), ("experience", "Experience"), ("skills", "Skills"), ("projects", "Projects")):
        chk(k in s, f"{lab} heading detected", f"Add a clear '{lab}' heading so ATS software can find the section.", 10)
    chk(meta["tables"] == 0, "No tables", "Tables can scramble ATS parsing; use plain lines or bullets.", 8)
    chk(meta["images"] == 0, "No images or graphics", "Remove images, icons and logos.", 6)
    chk(not any(unicodedata.category(ch) in ("So", "Co", "Cc", "Cs") and ch not in "\n\t" for ch in text), "No icon or symbol characters", "Replace icons and special symbols (✆ ✉ ★ etc.) with plain text.", 5)
    chk(250 <= words <= 1000, "Reasonable length (250-1000 words)", "Aim for 250-1000 words (roughly 1-2 pages).", 6)
    chk(len(parsed["skills"]) >= 6, "At least 6 recognised skills", "List more relevant technical skills.", 10)
    ats = round(100 * sum(c["weight"] for c in checks if c["pass"]) / sum(c["weight"] for c in checks))
    verbs = verb_hits / len(bullets) if bullets else 0
    if bullets and verbs < .5: recs.append("Start more bullet points with strong action verbs (built, automated, led).")
    if bullets and metrics == 0: recs.append("Consider adding a measurable result to some bullets if you have real numbers.")
    if not bullets: recs.append("Describe experience or projects in short bullet points.")
    scores = {"skills": min(100, len(parsed["skills"]) * 9), "experience": round(min(100, len(bullets) * 6 * .6 + verbs * 40)),
              "projects": min(100, len(bullet_items(s.get("projects", []))) * 12), "education": 90 if "education" in s else 30, "formatting": ats}
    overall = round(.25 * scores["skills"] + .25 * scores["experience"] + .15 * scores["projects"] + .1 * scores["education"] + .25 * ats)
    return {"overall_score": overall, "ats_score": ats, "scores": scores, "ats_checks": checks, "recommendations": recs,
            "stats": {"words": words, "bullets": len(bullets), "action_verb_bullets": verb_hits, "bullets_with_numbers": metrics},
            "note": "ATS score is an internal compatibility indicator, not a guarantee that any specific ATS will accept the resume."}

class AIError(Exception): pass
def llm(system, prompt, max_tokens=1000):
    """Calls Google's Gemini API (free tier, no credit card needed at signup). Set AI_API_KEY to a Gemini key from aistudio.google.com."""
    if not AI_KEY: raise AIError("AI is not configured. Set the AI_API_KEY environment variable.")
    import httpx
    url = f"https://generativelanguage.googleapis.com/v1beta/models/{AI_MODEL}:generateContent?key={AI_KEY}"
    try:
        r = httpx.post(url, timeout=60, json={
            "system_instruction": {"parts": [{"text": system}]},
            "contents": [{"role": "user", "parts": [{"text": prompt}]}],
            "generationConfig": {"maxOutputTokens": max_tokens}})
        r.raise_for_status()
        data = r.json()
        cand = data.get("candidates") or [None]
        if not cand[0] or not cand[0].get("content", {}).get("parts"):
            raise ValueError("empty response")  # e.g. blocked by safety filters
        return cand[0]["content"]["parts"][0]["text"]
    except Exception:
        raise AIError("AI is temporarily unavailable. Please try again.")
NO_FAKE = "Never invent experience, employers, degrees, certifications, skills or metrics. If a metric is missing, suggest adding one instead of making one up."

# ---------- resumes ----------
def own(c, rid, uid):
    r = c.execute("SELECT * FROM resumes WHERE id=? AND user_id=?", (rid, uid)).fetchone()
    if not r: raise HTTPException(404, "Resume not found.")
    return r

@app.post("/api/resumes/upload", status_code=201)
async def upload(file: UploadFile = File(...), uid=Depends(user)):
    data = await file.read()
    if len(data) > MAX_MB * 1024 * 1024: raise HTTPException(413, "Your file exceeds the maximum allowed size (5 MB).")
    try:
        text, meta = extract_text(file.filename or "", data)
    except HTTPException: raise
    except Exception:
        raise HTTPException(422, "We couldn't process this file. Try exporting it again as a PDF or DOCX.")
    if len(text) < 60: raise HTTPException(422, f"We could only read {len(text)} characters from this file. If it is a photo or scan, use a sharper, straight-on, higher-resolution image, or upload the PDF/DOCX.")
    parsed = parse(text); a = analyse(text, parsed, meta)
    c = db(); cur = c.execute("INSERT INTO resumes(user_id,filename,text,parsed,analysis) VALUES(?,?,?,?,?)",
                              (uid, Path(file.filename).name[:120], text, json.dumps(parsed), json.dumps(a))); c.commit()
    return {"id": cur.lastrowid}

@app.get("/api/resumes")
def list_resumes(uid=Depends(user)):
    rows = db().execute("SELECT id,filename,analysis,created FROM resumes WHERE user_id=? ORDER BY id DESC", (uid,)).fetchall()
    return [{"id": r["id"], "filename": r["filename"], "created": r["created"], "overall": json.loads(r["analysis"])["overall_score"],
             "ats": json.loads(r["analysis"])["ats_score"]} for r in rows]

@app.get("/api/resumes/{rid}")
def get_resume(rid: int, uid=Depends(user)):
    r = own(db(), rid, uid)
    return {"id": r["id"], "filename": r["filename"], "created": r["created"], "parsed": json.loads(r["parsed"]), "analysis": json.loads(r["analysis"])}

class Rename(BaseModel): filename: str = Field(min_length=1, max_length=120)
@app.patch("/api/resumes/{rid}")
def rename(rid: int, b: Rename, uid=Depends(user)):
    c = db(); own(c, rid, uid); c.execute("UPDATE resumes SET filename=? WHERE id=?", (b.filename, rid)); c.commit(); return {"ok": True}

@app.delete("/api/resumes/{rid}", status_code=204)
def delete(rid: int, uid=Depends(user)):
    c = db(); own(c, rid, uid); c.execute("DELETE FROM resumes WHERE id=?", (rid,)); c.commit()

@app.post("/api/resumes/{rid}/duplicate", status_code=201)
def duplicate(rid: int, uid=Depends(user)):
    c = db(); r = own(c, rid, uid)
    cur = c.execute("INSERT INTO resumes(user_id,filename,text,parsed,analysis) VALUES(?,?,?,?,?)", (uid, "Copy of " + r["filename"][:100], r["text"], r["parsed"], r["analysis"])); c.commit()
    return {"id": cur.lastrowid}

@app.get("/api/resumes/{a}/compare/{b}")
def compare(a: int, b: int, uid=Depends(user)):
    c = db(); ra, rb = own(c, a, uid), own(c, b, uid)
    A_, B_ = json.loads(ra["analysis"]), json.loads(rb["analysis"])
    sa, sb = set(json.loads(ra["parsed"])["skills"]), set(json.loads(rb["parsed"])["skills"])
    keys = ["overall_score", "ats_score"]
    return {"a": {"name": ra["filename"], **{k: A_[k] for k in keys}, "scores": A_["scores"]}, "b": {"name": rb["filename"], **{k: B_[k] for k in keys}, "scores": B_["scores"]},
            "only_in_a": sorted(sa - sb), "only_in_b": sorted(sb - sa), "shared": sorted(sa & sb)}

@app.post("/api/resumes/{rid}/ai-review")
def ai_review(rid: int, uid=Depends(user)):
    r = own(db(), rid, uid)
    try:
        out = llm("You are a careful resume reviewer. " + NO_FAKE, "Give 3 strengths and 3 weaknesses (short bullets) for this resume:\n\n" + r["text"][:6000])
    except AIError as e: raise HTTPException(503, str(e))
    return {"review": out}

# ---------- jobs ----------
class JobIn(BaseModel):
    title: str = Field(min_length=1, max_length=120); description: str = Field(min_length=40, max_length=20000)
@app.post("/api/jobs", status_code=201)
def add_job(b: JobIn, uid=Depends(user)):
    c = db(); cur = c.execute("INSERT INTO jobs(user_id,title,description) VALUES(?,?,?)", (uid, b.title, b.description)); c.commit(); return {"id": cur.lastrowid}
@app.get("/api/jobs")
def jobs(uid=Depends(user)):
    return [dict(r) for r in db().execute("SELECT id,title,created FROM jobs WHERE user_id=? ORDER BY id DESC", (uid,))]

STOP = set("the and for with you your our are will have has this that from their they who all any can able work team years year experience required preferred strong good knowledge skills ability using use well must should including etc plus role candidate looking join across about also more most such other into over than then them these those what when where which while within without would could being been each both some many much very good great nice bonus ideal position responsibilities requirements qualifications company job".split())
@app.post("/api/jobs/{jid}/match")
def match(jid: int, resume_id: int, uid=Depends(user)):
    c = db(); r = own(c, resume_id, uid)
    j = c.execute("SELECT * FROM jobs WHERE id=? AND user_id=?", (jid, uid)).fetchone()
    if not j: raise HTTPException(404, "Job not found.")
    need, have = set(find_skills(j["description"])), set(json.loads(r["parsed"])["skills"])
    kw = lambda t: {w.rstrip(".-") for w in re.findall(r"[a-z][a-z+#.-]{3,}", t.lower()) if w.rstrip(".-") not in STOP and len(w.rstrip(".-")) > 3}
    jk, rk = kw(j["description"]), kw(r["text"])
    jk_top = {w for w in jk if w not in need}
    skill_pct = len(need & have) / len(need) if need else None
    kw_pct = len(jk_top & rk) / len(jk_top) if jk_top else 0
    score = round(100 * (.7 * skill_pct + .3 * kw_pct)) if skill_pct is not None else round(100 * kw_pct)
    res = {"score": score, "matching_skills": sorted(need & have), "missing_skills": sorted(need - have),
           "missing_keywords": sorted(jk_top - rk)[:15], "matching_keywords": sorted(jk_top & rk)[:15],
           "confidence": "good" if len(need) >= 4 else "low - few recognisable skills in this job description",
           "note": "Keyword-based estimate. Missing items may be covered by different wording in your resume."}
    c.execute("INSERT INTO matches(resume_id,job_id,score,result) VALUES(?,?,?,?)", (resume_id, jid, score, json.dumps(res))); c.commit()
    return res

# ---------- AI ----------
class Improve(BaseModel): text: str = Field(min_length=3, max_length=2000); kind: str = "bullet"
@app.post("/api/ai/improve")
def improve(b: Improve, uid=Depends(user)):
    try: out = llm("You rewrite resume text. " + NO_FAKE + " Return only the rewritten text.", f"Improve this resume {b.kind}:\n{b.text}", 400)
    except AIError as e: raise HTTPException(503, str(e))
    c = db(); c.execute("INSERT INTO messages(user_id,role,content) VALUES(?,?,?)", (uid, "improve", b.text[:200])); c.commit()
    return {"improved": out.strip()}
class Chat(BaseModel): message: str = Field(min_length=1, max_length=2000); resume_id: int | None = None
@app.post("/api/ai/chat")
def chat(b: Chat, uid=Depends(user)):
    ctx = ""
    if b.resume_id: ctx = "\n\nUser's resume:\n" + own(db(), b.resume_id, uid)["text"][:5000]
    try: out = llm("You are a concise career assistant. " + NO_FAKE, b.message + ctx)
    except AIError as e: raise HTTPException(503, str(e))
    c = db(); c.executemany("INSERT INTO messages(user_id,role,content) VALUES(?,?,?)", [(uid, "user", b.message), (uid, "assistant", out)]); c.commit()
    return {"reply": out}

@app.get("/api/ai/history")
def chat_history(uid=Depends(user)):
    return [dict(r) for r in db().execute("SELECT role,content FROM messages WHERE user_id=? AND role IN ('user','assistant') ORDER BY id DESC LIMIT 40", (uid,))][::-1]

@app.delete("/api/auth/account", status_code=204)
def delete_account(resp: Response, uid=Depends(user)):
    c = db(); c.execute("DELETE FROM users WHERE id=?", (uid,)); c.commit(); resp.delete_cookie("session")

class Reset(BaseModel): email: str; new_password: str = Field(min_length=8, max_length=128); code: str
@app.post("/api/auth/reset-password")
def reset_password(b: Reset, request: Request):
    """Email isn't wired up, so resets are authorised by a server-side RESET_CODE the admin shares (e.g. with the user in class)."""
    limit("p" + request.client.host, 5)
    code = os.getenv("RESET_CODE")
    if not code or not hmac.compare_digest(b.code, code): raise HTTPException(403, "Reset code is invalid or password reset is not enabled.")
    c = db(); cur = c.execute("UPDATE users SET pw=? WHERE email=?", (hash_pw(b.new_password), b.email.lower())); c.commit()
    if not cur.rowcount: raise HTTPException(404, "No account with this email.")
    return {"ok": True}

# ---------- flow: matches, skill gaps, interview practice, application tracker ----------
@app.get("/api/matches")
def list_matches(uid=Depends(user)):
    rows = db().execute("SELECT m.id,m.score,m.result,m.created,j.title,r.filename,m.job_id,m.resume_id FROM matches m JOIN jobs j ON j.id=m.job_id "
                        "JOIN resumes r ON r.id=m.resume_id WHERE r.user_id=? ORDER BY m.id DESC", (uid,)).fetchall()
    return [{**{k: r[k] for k in ("id", "score", "created", "title", "filename", "job_id", "resume_id")}, "result": json.loads(r["result"])} for r in rows]

STATUSES = ["Saved", "Applied", "Interview", "Offer", "Rejected"]
class AppIn(BaseModel):
    company: str = Field(min_length=1, max_length=120); role: str = Field(min_length=1, max_length=120)
    status: str = "Saved"; notes: str = Field("", max_length=2000); resume_id: int | None = None; job_id: int | None = None
class AppPatch(BaseModel):
    status: str | None = None; notes: str | None = Field(None, max_length=2000)

def check_status(v):
    if v not in STATUSES: raise HTTPException(422, "Status must be one of: " + ", ".join(STATUSES))

@app.post("/api/applications", status_code=201)
def add_application(b: AppIn, uid=Depends(user)):
    check_status(b.status); c = db()
    if b.resume_id: own(c, b.resume_id, uid)
    if b.job_id and not c.execute("SELECT id FROM jobs WHERE id=? AND user_id=?", (b.job_id, uid)).fetchone(): raise HTTPException(404, "Job not found.")
    cur = c.execute("INSERT INTO applications(user_id,company,role,status,notes,resume_id,job_id) VALUES(?,?,?,?,?,?,?)",
                    (uid, b.company.strip(), b.role.strip(), b.status, b.notes, b.resume_id, b.job_id)); c.commit(); return {"id": cur.lastrowid}

@app.get("/api/applications")
def list_applications(uid=Depends(user)):
    return [dict(r) for r in db().execute("SELECT id,company,role,status,notes,resume_id,job_id,created FROM applications WHERE user_id=? ORDER BY id DESC", (uid,))]

@app.patch("/api/applications/{aid}")
def patch_application(aid: int, b: AppPatch, uid=Depends(user)):
    c = db()
    if not c.execute("SELECT id FROM applications WHERE id=? AND user_id=?", (aid, uid)).fetchone(): raise HTTPException(404, "Application not found.")
    if b.status is not None: check_status(b.status); c.execute("UPDATE applications SET status=? WHERE id=?", (b.status, aid))
    if b.notes is not None: c.execute("UPDATE applications SET notes=? WHERE id=?", (b.notes, aid))
    c.commit(); return {"ok": True}

@app.delete("/api/applications/{aid}", status_code=204)
def delete_application(aid: int, uid=Depends(user)):
    c = db(); cur = c.execute("DELETE FROM applications WHERE id=? AND user_id=?", (aid, uid)); c.commit()
    if not cur.rowcount: raise HTTPException(404, "Application not found.")

@app.get("/api/skill-gaps")
def skill_gaps(uid=Depends(user)):
    """Skills that job descriptions asked for that the matched resume lacked, ranked by how often they appear."""
    count, jobs_for = {}, {}
    for m in list_matches(uid):
        for sk in m["result"]["missing_skills"]:
            count[sk] = count.get(sk, 0) + 1; jobs_for.setdefault(sk, set()).add(m["title"])
    return [{"skill": k, "jobs": v, "job_titles": sorted(jobs_for[k]),
             "tip": f"If you have genuinely used {k}, add it to Skills or a project bullet. If not, a small project or course would close this gap; list it only once it is real."}
            for k, v in sorted(count.items(), key=lambda x: (-x[1], x[0]))]

class QIn(BaseModel): resume_id: int; job_id: int | None = None
@app.post("/api/interview/questions")
def interview_questions(b: QIn, uid=Depends(user)):
    c = db(); r = own(c, b.resume_id, uid); parsed = json.loads(r["parsed"]); gaps, title = [], ""
    if b.job_id:
        j = c.execute("SELECT * FROM jobs WHERE id=? AND user_id=?", (b.job_id, uid)).fetchone()
        if not j: raise HTTPException(404, "Job not found.")
        title = j["title"]; gaps = sorted(set(find_skills(j["description"])) - set(parsed["skills"]))
    qs = ["Tell me about yourself and why you are applying" + (f" for the {title} role." if title else ".")]
    qs += [f"You list {sk} on your resume. Describe a specific task where you used it and what the result was." for sk in parsed["skills"][:3]]
    bullets = [l for k in ("experience", "projects") for l in parsed["sections"].get(k, []) if len(l.split()) > 4][:2]
    qs += [f"Walk me through this item from your resume: \"{x[:110]}\". What was your role and what was the outcome?" for x in bullets]
    qs += [f"This role asks for {g}, which isn't on your resume. How would you get up to speed on it?" for g in gaps[:2]]
    qs += ["Describe a time you had a tight deadline. What did you do?", "Tell me about a disagreement in a team and how it ended."]
    source = "resume"
    if AI_KEY:
        try:
            out = llm("You are an interviewer. " + NO_FAKE, "Write 8 interview questions, one per line, no numbering, tailored to this resume" + (f" and the {title} role" if title else "") +
                      ". Only ask about things in the resume or typical for the role.\n\n" + r["text"][:5000], 600)
            ai = [re.sub(r"^[\d.\-\s]+", "", l).strip() for l in out.splitlines() if l.strip().endswith("?")]
            if len(ai) >= 4: qs, source = ai[:8], "ai"
        except AIError: pass
    return {"questions": qs, "source": source}

class FIn(BaseModel): question: str = Field(min_length=3, max_length=500); answer: str = Field(min_length=1, max_length=4000)
@app.post("/api/interview/feedback")
def interview_feedback(b: FIn, uid=Depends(user)):
    """Rule-based answer check (always on) plus optional AI comments."""
    a = b.answer.strip(); words = len(a.split()); tips, pts = [], 0
    def has(p): return re.search(p, a, re.I)
    checks = [(40 <= words <= 220, 25, "Aim for roughly 40-200 words: long enough to give detail, short enough to stay focused."),
              (has(r"\b(when|situation|project|class|internship|at work|team)\b"), 15, "Set the scene briefly: where and when this happened."),
              (has(r"\b(i )?(built|made|wrote|designed|led|fixed|tested|created|developed|configured|automated|handled|decided)\b"), 20, "Say what YOU did, using action verbs (built, fixed, tested)."),
              (has(r"\b(result|outcome|so that|which (led|meant|helped)|improved|reduced|saved|finished|learned)\b"), 25, "End with the result or what you learned."),
              (has(r"\d"), 15, "If you have a real number (time saved, users, marks), include it. Don't invent one.")]
    for ok, w, tip in checks:
        if ok: pts += w
        else: tips.append(tip)
    fb = {"score": pts, "tips": tips, "words": words, "ai_feedback": None}
    if AI_KEY:
        try: fb["ai_feedback"] = llm("You are a supportive interview coach. Be specific and brief. " + NO_FAKE, f"Question: {b.question}\nAnswer: {a}\nGive 3 short improvement points.", 400)
        except AIError: pass
    c = db(); c.execute("INSERT INTO interviews(user_id,question,answer,score,feedback) VALUES(?,?,?,?,?)", (uid, b.question, a, pts, json.dumps(fb))); c.commit()
    return fb

@app.get("/api/flow")
def flow(uid=Depends(user)):
    c = db()
    def n(q):
        row = c.execute(q, (uid,)).fetchone(); return list(row.values())[0] if PG else row[0]
    res = n("SELECT COUNT(*) FROM resumes WHERE user_id=?")
    ms = list_matches(uid); gaps = any(m["result"]["missing_skills"] for m in ms)
    steps = [("Upload", "upload", res > 0), ("Analyze", "dashboard", res > 0), ("Match job", "jobs", len(ms) > 0), ("Skill gaps", "gaps", gaps),
             ("Improve", "dashboard", n("SELECT COUNT(*) FROM messages WHERE user_id=? AND role='improve'") > 0), ("Interview", "interview", n("SELECT COUNT(*) FROM interviews WHERE user_id=?") > 0),
             ("Track", "apps", n("SELECT COUNT(*) FROM applications WHERE user_id=?") > 0)]
    return [{"label": l, "view": v, "done": d} for l, v, d in steps]

@app.get("/health")
def health():
    try: db().execute("SELECT 1").fetchone(); dbok = True
    except Exception: dbok = False
    return {"status": "ok" if dbok else "degraded", "database": "supabase-postgres" if PG else "sqlite", "database_ok": dbok, "ai_configured": bool(AI_KEY)}

@app.get("/")
def index(): return FileResponse(Path(__file__).parent / "static" / "index.html")
